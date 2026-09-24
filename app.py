# -*- coding: utf-8 -*-
"""嘟嘟换电统一看板 —— 同链接多菜单 + 筛选 + 导出

启动: python app.py
访问: http://127.0.0.1:8090/            (默认 用户协议)
      http://127.0.0.1:8090/?m=site     (网点)
      http://127.0.0.1:8090/?m=cabinet  (换电柜)
      http://127.0.0.1:8090/?m=battery  (电池)
"""
import csv
import io
import json
import os
import sqlite3
import sys
import threading
import time
from datetime import datetime
from urllib.parse import quote

from flask import Flask, Response, jsonify, render_template, request, send_file

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import schema_def as sd
from config import SERVICE_PORT, SQLITE_PATH

TEMP_DIR = os.path.join(BASE, "temp")
os.makedirs(TEMP_DIR, exist_ok=True)
XLSX_CACHE_TTL = 3600

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False

MENUS = ["agreement", "site", "cabinet", "battery"]
MENU_LABEL = sd.MENU_TITLE
COL_META = {m: sd.COLUMNS[m] for m in MENUS}
KEY_META = {m: {c["k"]: c for c in sd.COLUMNS[m] if c["t"] != "group"} for m in MENUS}

# 关键词搜索覆盖的列
KW_COLS = {
    "agreement": ["id", "user_id", "user_name", "user_phone", "cur_phone", "company_name",
                  "site_name", "battery_sn", "merchant_name", "guide_name", "business_name", "promoter_name"],
    "site": ["id", "name", "address", "detail_address", "merchant_name", "agency_name",
             "business_name", "distributor_name", "cabinet_sn", "income_owner"],
    "cabinet": ["sn", "name", "id", "site_name", "detail_address", "agency_name",
                "business_name", "imei", "iccid"],
    "battery": ["sn", "device_id", "cabinet_sn", "site_name", "model", "product", "location_desc"],
}

# 快捷筛选（key, 显示名）
QUICK = {
    "agreement": [("agreement_type", "协议类型"), ("status", "协议状态"), ("is_lowfreq", "是否低频"),
                  ("lowfreq_level", "低频档位"), ("is_first", "是否首次签约"), ("city", "城市"),
                  ("site_name", "签约网点"), ("merchant_name", "商户"), ("agency_name", "代理商"),
                  ("deposit_status", "押金状态")],
    "site": [("site_type", "网点类型"), ("site_status", "营业状态"), ("city", "城市"), ("area", "区域"),
             ("agency_name", "代理商"), ("merchant_name", "商户"), ("has_exchange", "是否可换电"),
             ("cabinet_status", "换电柜状态"), ("is_24h", "是否24小时")],
    "cabinet": [("online_status", "在线状态"), ("device_status", "设备状态"), ("city", "城市"),
                ("site_name", "所属网点"), ("agency_name", "代理商"), ("slot_total", "总仓位数"),
                ("rent_status", "租借状态")],
    "battery": [("online_status", "在线状态"), ("charge_status", "充电状态"), ("discharge_status", "放电状态"),
                ("city", "城市"), ("site_name", "所属网点"), ("cabinet_sn", "所在换电柜"), ("slot_type", "仓位类型")],
}

OPS = [("like", "包含"), ("eq", "等于"), ("ne", "不等于"), ("gt", "大于"), ("ge", "大于等于"),
       ("lt", "小于"), ("le", "小于等于"), ("between", "区间"), ("empty", "为空"), ("notempty", "不为空")]


def db():
    c = sqlite3.connect(SQLITE_PATH, timeout=60)
    c.row_factory = sqlite3.Row
    try:
        c.execute("PRAGMA cache_size=-64000")
        c.execute("PRAGMA mmap_size=268435456")
    except Exception:
        pass
    return c


def count_rows(c, menu, where, args, bare):
    """无筛选条件时直接读 sync 缓存的行数，避免全表扫描。"""
    if bare:
        try:
            r = c.execute("SELECT v FROM meta WHERE k=?", ("count_" + menu,)).fetchone()
            if r:
                return int(r[0])
        except Exception:
            pass
    return c.execute("SELECT COUNT(*) FROM %s WHERE %s" % (menu, where), args).fetchone()[0]


def is_real(t):
    return t in ("int", "num", "money")


def col_expr(k, t):
    return 'CAST("%s" AS REAL)' % k if is_real(t) else '"%s"' % k


def build_where(menu, conds, kw):
    """conds: [{"k":..,"op":..,"v":..}]  kw: 关键词"""
    keys = KEY_META[menu]
    where, args = [], []
    for c in conds or []:
        k = c.get("k")
        if k not in keys:
            continue
        op = c.get("op") or "like"
        t = keys[k]["t"]
        expr = col_expr(k, t)
        v = c.get("v")
        if op in ("empty", "notempty"):
            where.append("(\"%s\" IS NULL OR \"%s\" IN ('', '[]', '{}'))" % (k, k) if op == "empty"
                         else "(\"%s\" IS NOT NULL AND \"%s\" NOT IN ('', '[]', '{}'))" % (k, k))
            continue
        if v is None or (isinstance(v, str) and v.strip() == ""):
            continue
        if op == "between":
            vv = v if isinstance(v, (list, tuple)) else str(v).split(",")
            if len(vv) < 2:
                continue
            a, b = str(vv[0]).strip(), str(vv[1]).strip()
            if a == "" and b == "":
                continue
            if a != "":
                where.append("%s >= ?" % expr)
                args.append(float(a) if is_real(t) else a)
            if b != "":
                where.append("%s <= ?" % expr)
                args.append(float(b) if is_real(t) else b)
            continue
        if op == "like":
            where.append("\"%s\" LIKE ?" % k)
            args.append("%" + str(v).strip() + "%")
        else:
            sym = {"eq": "=", "ne": "<>", "gt": ">", "ge": ">=", "lt": "<", "le": "<="}[op]
            where.append("%s %s ?" % (expr, sym))
            try:
                args.append(float(v) if is_real(t) else str(v).strip())
            except Exception:
                args.append(str(v).strip())
    if kw:
        cols = [k for k in KW_COLS.get(menu, []) if k in keys]
        if cols:
            where.append("(" + " OR ".join("\"%s\" LIKE ?" % k for k in cols) + ")")
            args.extend(["%" + kw.strip() + "%"] * len(cols))
    return (" AND ".join(where) if where else "1=1"), args


def order_clause(menu, sort, order):
    keys = KEY_META[menu]
    if not sort or sort not in keys:
        return ""
    t = keys[sort]["t"]
    direction = "DESC" if str(order).lower() == "desc" else "ASC"
    return ' ORDER BY ("%s" IS NULL OR "%s" = \'\') ASC, %s %s' % (sort, sort, col_expr(sort, t), direction)


def parse_conds():
    raw = request.args.get("conds") or ""
    out = []
    if raw:
        try:
            d = json.loads(raw)
            if isinstance(d, list):
                out = d
        except Exception:
            out = []
    return out


def fmt_cell(k, t, v):
    if v is None:
        return ""
    v = str(v)
    if t == "money" and v not in ("", "-"):
        try:
            return ("%%.2f" % float(v))
        except Exception:
            return v
    return v


# ----------------------------------------------------------------- 页面
@app.route("/")
def index():
    menu = request.args.get("m") or "agreement"
    if menu not in MENUS:
        menu = "agreement"
    return render_template("index.html", menu=menu, menu_label=MENU_LABEL, menus=MENUS,
                           ops=OPS, quick={m: QUICK[m] for m in MENUS})


@app.route("/api/meta")
def api_meta():
    out = {}
    for m in MENUS:
        out[m] = {"label": MENU_LABEL[m], "cols": COL_META[m], "keys": sd.DATA_KEYS[m],
                  "quick": [{"k": k, "n": n} for k, n in QUICK[m]]}
    return jsonify({"menus": MENUS, "meta": out, "ops": OPS})


@app.route("/api/count")
def api_count():
    menu = request.args.get("m")
    if menu not in MENUS:
        return jsonify({"error": "bad menu"}), 400
    conds, kw = parse_conds(), request.args.get("kw")
    where, args = build_where(menu, conds, kw)
    bare = not conds and not kw
    c = db()
    try:
        n = count_rows(c, menu, where, args, bare)
    finally:
        c.close()
    return jsonify({"total": n})


@app.route("/api/rows")
def api_rows():
    menu = request.args.get("m")
    if menu not in MENUS:
        return jsonify({"error": "bad menu"}), 400
    page = max(int(request.args.get("page") or 1), 1)
    size = min(max(int(request.args.get("size") or 50), 1), 500)
    where, args = build_where(menu, parse_conds(), request.args.get("kw"))
    keys = sd.DATA_KEYS[menu]
    sql = 'SELECT %s FROM %s WHERE %s%s LIMIT ? OFFSET ?' % (
        ", ".join('"%s"' % k for k in keys), menu, where,
        order_clause(menu, request.args.get("sort"), request.args.get("order")))
    c = db()
    t0 = time.time()
    try:
        total = count_rows(c, menu, where, args, not parse_conds() and not request.args.get("kw"))
        rows = c.execute(sql, args + [size, (page - 1) * size]).fetchall()
    finally:
        c.close()
    meta = KEY_META[menu]
    data = [[fmt_cell(k, meta[k]["t"], r[i]) for i, k in enumerate(keys)] for r in rows]
    return jsonify({"total": total, "page": page, "size": size, "keys": keys, "rows": data,
                    "ms": int((time.time() - t0) * 1000)})


@app.route("/api/options")
def api_options():
    menu = request.args.get("m")
    k = request.args.get("k")
    if menu not in MENUS or k not in KEY_META[menu]:
        return jsonify({"error": "bad params"}), 400
    c = db()
    try:
        rows = c.execute('SELECT "%s" v, COUNT(*) n FROM %s WHERE "%s" NOT IN (\'\', \'[]\', \'{}\') '
                         'GROUP BY "%s" ORDER BY n DESC LIMIT 300' % (k, menu, k, k)).fetchall()
        total = c.execute("SELECT COUNT(*) FROM %s" % menu).fetchone()[0]
    finally:
        c.close()
    return jsonify({"total": total, "options": [{"v": r["v"], "n": r["n"]} for r in rows]})


# ----------------------------------------------------------------- 导出
def iter_rows(menu, conds, kw, fn, sort=None, order=None):
    where, args = build_where(menu, conds, kw)
    keys = sd.DATA_KEYS[menu]
    fn(keys)
    c = db()
    try:
        cur = c.execute('SELECT %s FROM %s WHERE %s%s' % (
            ", ".join('"%s"' % k for k in keys), menu, where,
            order_clause(menu, sort, order)), args)
        meta = KEY_META[menu]
        while True:
            chunk = cur.fetchmany(3000)
            if not chunk:
                break
            for r in chunk:
                yield [fmt_cell(k, meta[k]["t"], r[i]) for i, k in enumerate(keys)]
    finally:
        c.close()


def csv_stream(menu, conds, kw, sort=None, order=None):
    keys = sd.DATA_KEYS[menu]
    labels = [KEY_META[menu][k]["n"] for k in keys]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(labels)
    yield "\ufeff" + buf.getvalue()
    buf.seek(0)
    buf.truncate(0)
    n = 0
    for row in iter_rows(menu, conds, kw, lambda k: None, sort, order):
        w.writerow(row)
        n += 1
        if n % 3000 == 0:
            yield buf.getvalue()
            buf.seek(0)
            buf.truncate(0)
    tail = buf.getvalue()
    if tail:
        yield tail


def cleanup_temp(max_age=XLSX_CACHE_TTL):
    now = time.time()
    try:
        for f in os.listdir(TEMP_DIR):
            p = os.path.join(TEMP_DIR, f)
            if os.path.isfile(p) and f.endswith(".xlsx") and now - os.path.getmtime(p) > max_age:
                os.remove(p)
    except Exception:
        pass


@app.route("/api/export")
def api_export():
    menu = request.args.get("m")
    if menu not in MENUS:
        return jsonify({"error": "bad menu"}), 400
    fmt = (request.args.get("fmt") or "csv").lower()
    conds, kw = parse_conds(), request.args.get("kw")
    sort, order = request.args.get("sort"), request.args.get("order")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = "%s_%s" % (MENU_LABEL[menu], stamp)

    if fmt == "csv":
        resp = Response(csv_stream(menu, conds, kw, sort, order), mimetype="text/csv; charset=utf-8")
        resp.headers["Content-Disposition"] = "attachment; filename*=UTF-8''%s.csv" % quote(name)
        return resp

    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    keys = sd.DATA_KEYS[menu]
    meta = KEY_META[menu]
    path = os.path.join(TEMP_DIR, "%s.xlsx" % name)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=MENU_LABEL[menu][:30])
    ws.append([meta[k]["n"] for k in keys])
    for i, k in enumerate(keys, 1):
        n = meta[k]["n"]
        width = 12 if meta[k]["t"] in ("int", "num", "money") else 18
        if len(n) * 2 > width:
            width = min(len(n) * 2, 30)
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in iter_rows(menu, conds, kw, lambda k: None, sort, order):
        ws.append(row)
    wb.save(path)
    threading.Thread(target=cleanup_temp, daemon=True).start()
    return send_file(path, as_attachment=True, download_name="%s.xlsx" % name)


if __name__ == "__main__":
    print("统一看板启动: http://127.0.0.1:%d/  (数据库 %s)" % (SERVICE_PORT, SQLITE_PATH))
    app.run(host="0.0.0.0", port=SERVICE_PORT, threaded=True, debug=False)
